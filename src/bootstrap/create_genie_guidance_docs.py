"""Build the five seller-guidance DOCX files consumed by the Genie Agent.

The source Markdown stays human-reviewable in ``src/bootstrap/documents``.  This
script produces a deliberately simple Word representation because Genie volume
retrieval depends on clean document structure rather than decorative layout.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from guidance_documents import SOURCE_MARKDOWN_NAMES


SOURCE_NAMES = SOURCE_MARKDOWN_NAMES


def _set_font(style, *, size: float, bold: bool = False) -> None:
    font = style.font
    font.name = "Arial"
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor(0, 0, 0)
    style.element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style.element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    _set_font(document.styles["Normal"], size=11)
    normal = document.styles["Normal"]
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.08

    _set_font(document.styles["Title"], size=20, bold=True)
    title = document.styles["Title"]
    title.paragraph_format.space_after = Pt(14)
    title_properties = title.element.get_or_add_pPr()
    for tag in ("w:pBdr", "w:shd"):
        for node in title_properties.findall(qn(tag)):
            title_properties.remove(node)

    _set_font(document.styles["Heading 1"], size=14, bold=True)

    if "Source Note" not in document.styles:
        source_note = document.styles.add_style("Source Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source_note = document.styles["Source Note"]
    _set_font(source_note, size=9)
    source_note.paragraph_format.space_before = Pt(14)


def _strip_inline_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def build_docx(source: Path, output: Path) -> None:
    document = Document()
    _configure_document(document)
    lines = source.read_text(encoding="utf-8").splitlines()
    title_written = False

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# ") and not title_written:
            document.add_paragraph(_strip_inline_markdown(line[2:]), style="Title")
            title_written = True
            continue
        if line.startswith("## "):
            document.add_paragraph(_strip_inline_markdown(line[3:]), style="Heading 1")
            continue
        if line.startswith("- "):
            document.add_paragraph(_strip_inline_markdown(line[2:]), style="List Bullet")
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            document.add_paragraph(_strip_inline_markdown(numbered.group(2)), style="List Number")
            continue
        document.add_paragraph(_strip_inline_markdown(line))

    if not title_written:
        raise ValueError(f"{source} must start with a level-one title")
    document.add_paragraph(
        "Seller guidance source for the Agentic CPQ demo. Verify commercial terms in governed data before quoting.",
        style="Source Note",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source-dir",
        type=Path,
        default=Path(__file__).with_name("documents"),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).with_name("genie_guidance_docs"),
    )
    args = parser.parse_args()

    for name in SOURCE_NAMES:
        source = args.source_dir / name
        if not source.is_file():
            raise FileNotFoundError(source)
        output = args.output_dir / f"{source.stem}.docx"
        build_docx(source, output)
        print(output)


if __name__ == "__main__":
    main()
